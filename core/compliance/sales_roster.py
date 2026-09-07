"""Parser satu-satunya untuk "Update Sales Telemarketing …xlsx" (database sales).

Dipisahkan dari ``api.sales_lookup`` supaya WORKER bisa ikut membacanya tanpa menarik
FastAPI: worker perlu kolom ``NAME ONLINE`` untuk menentukan panggilan mana milik agent
yang di-assign TMS (lihat ``compliance.call_ownership``). ``api.sales_lookup`` tetap
memegang pengambilan file dari MinIO, cache-nya, dan seluruh helper cakupan/hierarki —
yang pindah ke sini HANYA pembacaan sheet-nya, supaya aturan kolomnya tidak pernah ada
dua versi yang bisa berbeda.

Seluruh kolom identitas dibaca lewat ``person()``: placeholder roster ("-", "0",
"00000000", "#N/A", …) menjadi string kosong, dan baris padding ber-USER ID "0"
tidak menjadi entri sama sekali. Tanpa itu dropdown "Semua AM"/"Semua TL" di halaman
Results memunculkan satu opsi bernama ``-`` dan agen MUTASI tampak masih punya
atasan — lihat ``tests/test_sales_roster_placeholders.py``.

Tata letak sheet (dicocokkan lewat HEADER dulu, baru posisi kolom tetap):
  A USER ID | C NIP BARU | D NAME | E NAME ONLINE | F DEDICATED | H NIP TL |
  I NAMA TL | J NIP TLM (= NIP AM) | K NAMA AM | O JOIN POSISI (DD/MM/YYYY)
"""
import io
from datetime import date, datetime

from openpyxl import load_workbook

# JOIN POSISI cells are usually datetimes; when a string, they look like "31/10/2017".
JOIN_FORMATS = ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d")

# Sales-marketing hierarchy columns: matched by header first, else fixed position.
# Column I = "NAMA TL" (index 8, Team Leader), column K = "NAMA AM" (index 10, Area
# Manager). See the "Update Sales Telemarketing …xlsx" layout.
_TL_HEADER = "nama tl"
_AM_HEADER = "nama am"
_TL_FALLBACK_IDX = 8   # column I
_AM_FALLBACK_IDX = 10  # column K
# Sales-agent (Team Leader) login scoping: NIP TL = column H (idx 7), the NIP of the
# agent's team leader; DEDICATED = column F (idx 5), the campaign the agent handles.
_NIP_TL_HEADER = "nip tl"
_DEDICATED_HEADER = "dedicated"
_NIP_TL_FALLBACK_IDX = 7  # column H
_DEDICATED_FALLBACK_IDX = 5  # column F
# Area Manager login scoping: NIP AM = column J (idx 9), the NIP of the agent's area
# manager (one level above the Team Leader), paired with column K = NAMA AM. In the
# "Update Sales Telemarketing …" sheet column J is labelled "NIP TLM", so accept both
# that header and "NIP AM"; fall back to the fixed column J when neither is present.
_NIP_AM_HEADERS = ("nip am", "nip tlm")
_NIP_AM_FALLBACK_IDX = 9  # column J
# QC (agent) login scoping: NIP BARU = column C (idx 2), the agent's own NIP.
_NIP_BARU_HEADER = "nip baru"
_NIP_BARU_FALLBACK_IDX = 2  # column C
# NAME ONLINE = column E (idx 4), nama panggilan yang dipakai agent saat menelepon
# nasabah — beda dari kolom D "NAME" yang berisi nama lengkap karyawan. Ditampilkan
# di Agent Error Summary di sebelah kanan Agent Name, dan dipakai
# ``compliance.call_ownership`` untuk mengenali pemilik sebuah panggilan.
_NAME_ONLINE_HEADER = "name online"
_NAME_ONLINE_FALLBACK_IDX = 4  # column E


def norm(value) -> str:
    """Trimmed string for header/cell matching; an integer-valued float loses its
    ``.0`` so a numeric USER ID like ``801.0`` matches ``"801"`` ('' for None)."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


# Kolom roster jarang dikosongkan — yang tidak ada isinya diberi PLACEHOLDER:
# 12 baris padding di akhir sheet memakai USER ID "0" / NIP "00000000" / nama "-",
# dan agen berstatus MUTASI kehilangan atasannya dengan cara yang sama (NIP TL &
# NIP TLM "0", NAMA TL & NAMA AM "-"). Placeholder BUKAN identitas orang: dibiarkan
# apa adanya, NIP AM "0" + NAMA AM "-" menjadi satu opsi "-" di dropdown Semua AM
# (idem Semua TL) dan baris padding menjadi satu "agent" hantu.
_PLACEHOLDERS = {"-", "--", "n/a", "na", "none", "null", "#n/a", "#ref!"}


def person(value) -> str:
    """``norm``, tapi placeholder roster ("-", "0", "00000000", …) jadi ''."""
    s = norm(value)
    if not s or s.casefold() in _PLACEHOLDERS:
        return ""
    return "" if set(s) == {"0"} else s


def to_date(value, formats) -> "date | None":
    """Coerce a datetime/date/string cell into a ``date`` (None if unparseable)."""
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    s = str(value).strip()
    if not s:
        return None
    for fmt in formats:
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def _pick(norm_header, header_names, fallback_idx):
    """Indeks kolom: cocokkan HEADER dulu, baru jatuh ke posisi kolom tetap."""
    if isinstance(header_names, str):
        header_names = (header_names,)
    idx = next((i for i, h in enumerate(norm_header) if h in header_names), None)
    if idx is None and fallback_idx is not None and len(norm_header) > fallback_idx:
        idx = fallback_idx
    return idx


def parse_roster(data: bytes) -> dict:
    """``{USER ID (casefold) -> {"name", "name_online", "join_date", "team_leader",
    "area_manager", "nip_tl", "nip_am", "dedicated", "nip_baru"}}`` dari isi xlsx.

    Mengembalikan ``{}`` bila file-nya tidak bisa dibaca/di-parse — pemanggil
    memperlakukannya sama dengan "tidak ada database sales aktif".
    """
    mapping: dict = {}
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        header = next(rows, None)
        if header:
            nh = [norm(h).casefold() for h in header]
            uid_i = _pick(nh, ("user id", "user_id", "userid"), None)
            name_i = _pick(nh, ("name",), None)
            nameon_i = _pick(nh, (_NAME_ONLINE_HEADER,), _NAME_ONLINE_FALLBACK_IDX)
            join_i = next((i for i, h in enumerate(nh) if h.startswith("join posisi")), None)
            tl_i = _pick(nh, (_TL_HEADER,), _TL_FALLBACK_IDX)
            am_i = _pick(nh, (_AM_HEADER,), _AM_FALLBACK_IDX)
            niptl_i = _pick(nh, (_NIP_TL_HEADER,), _NIP_TL_FALLBACK_IDX)
            ded_i = _pick(nh, (_DEDICATED_HEADER,), _DEDICATED_FALLBACK_IDX)
            nipam_i = _pick(nh, _NIP_AM_HEADERS, _NIP_AM_FALLBACK_IDX)
            nipbaru_i = _pick(nh, (_NIP_BARU_HEADER,), _NIP_BARU_FALLBACK_IDX)
            if uid_i is not None:
                # SEMUA kolom identitas dibaca lewat person(), bukan norm():
                # placeholder roster harus jadi kosong SEBELUM tersimpan, supaya tidak
                # ada konsumen (dropdown filter, hierarki Statistics, scoping login,
                # compliance.call_ownership) yang perlu tahu soal "-" dan "0".
                cell = lambda r, i: person(r[i]) if (i is not None and i < len(r)) else ""
                for r in rows:
                    if not r or uid_i >= len(r):
                        continue
                    uid = person(r[uid_i])
                    if not uid:
                        continue  # baris padding "0" — bukan agent
                    join_date = (
                        to_date(r[join_i], JOIN_FORMATS)
                        if (join_i is not None and join_i < len(r))
                        else None
                    )
                    mapping[uid.casefold()] = {
                        "name": cell(r, name_i) or None,
                        "name_online": cell(r, nameon_i) or None,
                        "join_date": join_date,
                        "team_leader": cell(r, tl_i) or None,
                        "area_manager": cell(r, am_i) or None,
                        "nip_tl": cell(r, niptl_i) or None,
                        "nip_am": cell(r, nipam_i) or None,
                        "dedicated": cell(r, ded_i) or None,
                        "nip_baru": cell(r, nipbaru_i) or None,
                    }
        wb.close()
    except Exception:
        return {}
    return mapping
